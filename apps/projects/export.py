import base64
import io
import os
import re
from datetime import date

from django.http import HttpResponse
from django.utils.translation import gettext as _
from django.utils.translation import ngettext
from PIL import Image, ImageOps

from apps.administration.models import SystemSettings
from apps.projects.models import ProjectAttribute, attribute_user_names
from apps.projects.models import Item, ProjectMembership, ProjectVersion, Section


def export_project(project, export_format, options):
    sys_settings = SystemSettings.load()
    options['_org_name'] = sys_settings.org_name or ''
    options['_year'] = date.today().year
    options['_user_names'] = attribute_user_names(project)
    sections = list(Section.objects.filter(project=project).select_related('parent').order_by('order'))
    items = Item.objects.filter(project=project).select_related(
        'section', 'assignee', 'created_by'
    ).prefetch_related(
        'item_tags__tag', 'comments__author', 'attribute_values__attribute',
        'outgoing_links__target_item__project',
        'incoming_links__source_item__project',
    )
    items = items.order_by('section__order', 'order')

    items_by_section = {}
    for item in items:
        items_by_section.setdefault(item.section_id, []).append(item)

    if export_format == 'docx':
        return _export_docx(project, sections, items_by_section, options)
    return _export_pdf(project, sections, items_by_section, options)


IMAGE_DPI = 150
IMAGE_MAX_HEIGHT_CM = 14
IMAGE_MAX_HEIGHT_EMU = IMAGE_MAX_HEIGHT_CM * 360000
PASSTHROUGH_FORMATS = {'JPEG': 'image/jpeg', 'PNG': 'image/png', 'GIF': 'image/gif'}
PASSTHROUGH_MODES = {'RGB', 'RGBA', 'L', 'P'}


def _prepare_image(image_field):
    if not image_field:
        return None
    try:
        path = image_field.path
        if not os.path.exists(path):
            return None
        with Image.open(path) as source:
            source_format = source.format
            source_mode = source.mode
            orientation = source.getexif().get(0x0112, 1)
            if orientation == 1 and source_format in PASSTHROUGH_FORMATS and source_mode in PASSTHROUGH_MODES:
                with open(path, 'rb') as f:
                    return f.read(), PASSTHROUGH_FORMATS[source_format], source.width, source.height
            image = ImageOps.exif_transpose(source)
            if image is None:
                image = source.copy()
        if image.mode == 'CMYK':
            image = image.convert('RGB')
        elif image.mode in ('I;16', 'I;16B', 'I;16L', 'I;16N', 'I', 'F'):
            image = image.convert('I').point(lambda value: value / 256).convert('L')
        elif image.mode in ('LA', 'PA', 'P', 'RGBa', 'La'):
            image = image.convert('RGBA')
        elif image.mode not in ('RGB', 'RGBA', 'L'):
            image = image.convert('RGB')
        buffer = io.BytesIO()
        if image.mode == 'RGBA':
            image.save(buffer, format='PNG')
            mime = 'image/png'
        else:
            image.save(buffer, format='JPEG', quality=90)
            mime = 'image/jpeg'
        return buffer.getvalue(), mime, image.width, image.height
    except Exception:
        return None


def _fit_image_width_emu(width_px, height_px, max_width_emu, max_height_emu):
    native_width = width_px * 914400 / IMAGE_DPI
    native_height = height_px * 914400 / IMAGE_DPI
    scale = min(1.0, max_width_emu / native_width, max_height_emu / native_height)
    return int(native_width * scale)


def _get_item_links(item):
    outgoing = []
    for link in item.outgoing_links.all():
        outgoing.append({
            'type': link.get_link_type_display(),
            'item_id': link.target_item.item_id,
            'item_pk': link.target_item.pk,
            'project': link.target_item.project.name,
            'project_id': link.target_item.project_id,
        })
    incoming = []
    for link in item.incoming_links.all():
        incoming.append({
            'type': link.get_link_type_display(),
            'item_id': link.source_item.item_id,
            'item_pk': link.source_item.pk,
            'project': link.source_item.project.name,
            'project_id': link.source_item.project_id,
        })
    return outgoing, incoming


def _get_export_attributes(project, options):
    if not options.get('include_attributes'):
        return []
    attr_ids = options.get('export_attr_ids', [])
    if not attr_ids:
        return []
    return list(
        ProjectAttribute.objects.filter(
            project=project, pk__in=attr_ids, is_builtin=False
        ).order_by('order')
    )


def _attribute_lines(item, export_attrs, user_names):
    values = {av.attribute_id: av.value for av in item.attribute_values.all()}
    lines = []
    for attr in export_attrs:
        value = values.get(attr.pk)
        if value is None or value == '':
            continue
        lines.append((attr.name, attr.display_value(value, user_names)))
    return lines


def _shows_empty_note(section, items_by_section, parent_ids):
    return not items_by_section.get(section.pk) and section.pk not in parent_ids


def _get_authors(project):
    authors = [str(project.created_by)]
    for m in ProjectMembership.objects.filter(project=project).select_related('user').order_by('added_at'):
        if m.user_id == project.created_by_id:
            continue
        authors.append(str(m.user))
    return authors


def _get_latest_version_number(project):
    version = ProjectVersion.objects.filter(project=project).order_by('-changed_at').first()
    return version.version_number if version else None


def _restricted_url_fetcher():
    from weasyprint import URLFetcher
    return URLFetcher(allowed_protocols=('data',))


STATUS_COLORS = {
    'draft': (0x47, 0x55, 0x69),
    'in-review': (0x92, 0x40, 0x0e),
    'approved': (0x16, 0x65, 0x34),
    'rejected': (0x99, 0x1b, 0x1b),
}
PRIORITY_COLORS = {
    'high': (0xdc, 0x26, 0x26),
    'medium': (0xd9, 0x77, 0x06),
    'low': (0x16, 0xa3, 0x4a),
}
PAGE_FIELDS = re.compile(r'%\((page|pages)\)s')
PAGE_FIELD_CODES = {'page': 'PAGE', 'pages': 'NUMPAGES'}
VALID_PAPER_SIZES = {'a4', 'letter', 'legal', 'a3'}
VALID_ORIENTATIONS = {'portrait', 'landscape'}
PAPER_SIZES_MM = {'a4': (210, 297), 'letter': (215.9, 279.4), 'legal': (215.9, 355.6), 'a3': (297, 420)}


def _page_label_parts():
    text = _('Page %(page)s of %(pages)s')
    parts = []
    position = 0
    for match in PAGE_FIELDS.finditer(text):
        parts.append(('text', text[position:match.start()]))
        parts.append(('field', match.group(1)))
        position = match.end()
    parts.append(('text', text[position:]))
    return [part for part in parts if part != ('text', '')]


def _css_page_label():
    tokens = []
    for kind, value in _page_label_parts():
        if kind == 'field':
            tokens.append(f'counter({value})')
        else:
            escaped = value.replace('\\', '\\\\').replace('"', '\\"')
            tokens.append(f'"{escaped}"')
    return ' '.join(tokens)


def _authors_line(authors):
    return ngettext('Author: %(names)s', 'Authors: %(names)s', len(authors)) % {'names': ', '.join(authors)}


def page_setup(options):
    paper = options.get('paper_size', 'a4').lower()
    if paper not in VALID_PAPER_SIZES:
        paper = 'a4'
    orientation = options.get('orientation', 'portrait').lower()
    if orientation not in VALID_ORIENTATIONS:
        orientation = 'portrait'
    return paper, orientation


def _export_pdf(project, sections, items_by_section, options):
    try:
        from weasyprint import HTML
    except ImportError:
        return _export_pdf_fallback(project, sections, items_by_section, options)

    html_content = _build_export_html(project, sections, items_by_section, options)

    pdf_buffer = io.BytesIO()
    HTML(string=html_content, url_fetcher=_restricted_url_fetcher()).write_pdf(pdf_buffer)
    pdf_buffer.seek(0)

    response = HttpResponse(pdf_buffer.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{project.slug}.pdf"'
    return response


def _build_export_html(project, sections, items_by_section, options):
    paper, orientation = page_setup(options)
    paper = paper.upper()
    org_name = _esc(options.get('_org_name', ''))
    year = options.get('_year', date.today().year)
    footer_left = f'{org_name} © {year}' if org_name else f'© {year}'
    user_names = options.get('_user_names', {})
    parent_ids = {s.parent_id for s in sections if s.parent_id}
    export_attrs = _get_export_attributes(project, options)
    include_assignee = bool(options.get('include_assignee'))
    page_label = _css_page_label()

    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
@page {{
  size: {paper} {orientation};
  margin: 1.5cm 1.5cm 2cm 1.5cm;
  @bottom-left {{
    content: "{footer_left}";
    font-size: 8pt;
    color: #999;
  }}
  @bottom-right {{
    content: {page_label};
    font-size: 8pt;
    color: #999;
  }}
}}
@page title-page {{
  @bottom-left {{ content: none; }}
  @bottom-right {{ content: none; }}
}}
@page toc-page {{
  @bottom-left {{ content: none; }}
  @bottom-right {{ content: none; }}
}}
body {{ font-family: sans-serif; font-size: 10pt; line-height: 1.4; color: #222; }}
h1 {{ color: #284E8A; border-bottom: 2px solid #284E8A; padding-bottom: 6px; font-size: 18pt; margin-top: 0; }}
h2 {{ color: #333; margin-top: 18px; font-size: 14pt; page-break-after: avoid; }}
h3 {{ color: #555; font-size: 12pt; page-break-after: avoid; }}
h4 {{ color: #555; font-size: 11pt; page-break-after: avoid; }}
.title-page {{
  page: title-page;
  page-break-after: always;
  padding-top: 10cm;
  text-align: center;
}}
.title-page h1 {{ font-size: 28pt; border: none; margin: 0; }}
.title-page .description {{ font-size: 12pt; color: #555; margin: 8px 2cm 0; }}
.title-page .version {{ font-size: 11pt; color: #666; margin-top: 16px; }}
.title-page .authors {{ font-size: 10pt; color: #666; margin-top: 12px; }}
.toc-wrapper {{ page: toc-page; page-break-after: always; }}
table {{ width: 100%; border-collapse: collapse; margin: 8px 0; font-size: 9pt; }}
th {{ background: #f1f5f9; border: 1px solid #d1d5db; padding: 4px 6px; text-align: left; font-weight: 600; }}
td {{ border: 1px solid #d1d5db; padding: 4px 6px; vertical-align: top; }}
tr {{ page-break-inside: avoid; }}
.badge {{ display: inline-block; padding: 1px 6px; border-radius: 3px; font-size: 8pt; }}
.status-draft {{ background: #e2e8f0; color: #475569; }}
.status-in-review {{ background: #fef3c7; color: #92400e; }}
.status-approved {{ background: #dcfce7; color: #166534; }}
.status-rejected {{ background: #fee2e2; color: #991b1b; }}
.priority-high {{ color: #dc2626; font-weight: bold; }}
.priority-medium {{ color: #d97706; }}
.priority-low {{ color: #16a34a; }}
.comment {{ background: #f8f9fa; padding: 6px 8px; margin: 3px 0; border-left: 3px solid #284E8A; font-size: 8pt; }}
.item-image {{ max-width: 100%; max-height: {IMAGE_MAX_HEIGHT_CM}cm; width: auto; height: auto; margin: 4px 0; }}
.image-caption {{ font-size: 8pt; color: #666; font-style: italic; margin-bottom: 4px; }}
.attrs {{ font-size: 8pt; color: #555; margin: 4px 0; }}
.attrs strong {{ color: #333; font-weight: 600; }}
.attrs .sep {{ color: #999; margin: 0 4px; }}
.links-section {{ font-size: 8pt; color: #555; margin: 4px 0; }}
.links-section strong {{ color: #333; }}
.item-text table {{ width: auto; margin: 6px 0; font-size: 8pt; border-collapse: collapse; }}
.item-text th {{ background: #e2e8f0; border: 1px solid #d1d5db; padding: 3px 5px; }}
.item-text td {{ border: 1px solid #d1d5db; padding: 3px 5px; }}
.item-text p {{ margin: 2px 0; }}
.toc {{ list-style: none; padding: 0; }}
.toc li {{ margin: 4px 0; padding: 2px 0; }}
.toc a {{ text-decoration: none; color: #284E8A; }}
.toc a:hover {{ text-decoration: underline; }}
</style></head><body>
<div class="title-page">
<h1>{_esc(project.name)}</h1>"""

    if project.description:
        html += f'<div class="description">{_esc(project.description)}</div>'

    version_number = _get_latest_version_number(project)
    if version_number:
        html += '<div class="version">' + _esc(_('Version %(number)s') % {'number': version_number}) + '</div>'

    if options.get('include_authors'):
        authors = _get_authors(project)
        if authors:
            html += '<div class="authors">' + _esc(_authors_line(authors)) + '</div>'

    html += '</div>'

    if options.get('include_toc'):
        html += '<div class="toc-wrapper"><h2 style="text-align:center">' + _esc(_('Table of Contents')) + '</h2><ul class="toc">'
        for section in sections:
            anchor = f'section-{section.pk}'
            depth = section.number.count('.')
            html += f'<li style="padding-left:{depth * 20}px"><a href="#{anchor}">{_esc(section.number)} {_esc(section.title)}</a></li>'
        html += '</ul></div>'

    for section in sections:
        anchor = f'section-{section.pk}'
        depth = section.number.count('.')
        tag = 'h2' if depth == 0 else ('h3' if depth == 1 else 'h4')
        html += f'<{tag} id="{anchor}">{_esc(section.number)} {_esc(section.title)}</{tag}>'

        section_items = items_by_section.get(section.pk, [])
        if not section_items:
            if _shows_empty_note(section, items_by_section, parent_ids):
                html += '<p><em>' + _esc(_('No items in this section.')) + '</em></p>'
            continue

        html += (
            f'<table><thead><tr><th style="width:70px">{_esc(_("ID"))}</th><th>{_esc(_("Text"))}</th>'
            f'<th style="width:70px">{_esc(_("Status"))}</th><th style="width:60px">{_esc(_("Priority"))}</th>'
        )
        if include_assignee:
            html += f'<th style="width:90px">{_esc(_("Assignee"))}</th>'
        html += '</tr></thead><tbody>'

        for item in section_items:
            status_class = f'status-{item.status}'
            priority_class = f'priority-{item.priority}'
            item_anchor = f'item-{item.pk}'

            text_content = '<div class="item-text">' + (item.text or '') + '</div>'

            prepared = _prepare_image(item.image)
            if prepared:
                data, mime, _width, _height = prepared
                img_data = f'data:{mime};base64,{base64.b64encode(data).decode("ascii")}'
                text_content += f'<br><img class="item-image" src="{img_data}">'
                if item.image_caption:
                    text_content += f'<div class="image-caption">{_esc(item.image_caption)}</div>'

            attr_lines = _attribute_lines(item, export_attrs, user_names)
            if attr_lines:
                text_content += '<div class="attrs">' + '<span class="sep">·</span>'.join(
                    f'<strong>{_esc(name)}:</strong> {_esc(value)}' for name, value in attr_lines
                ) + '</div>'

            if options.get('include_links'):
                outgoing, incoming = _get_item_links(item)
                if outgoing or incoming:
                    text_content += '<div class="links-section">'
                    for link in outgoing:
                        if link['project_id'] == project.pk:
                            item_ref = f'<a href="#item-{link["item_pk"]}" style="color:#284E8A"><strong>{_esc(link["item_id"])}</strong></a>'
                        else:
                            item_ref = f'<strong>{_esc(link["item_id"])}</strong> ({_esc(link["project"])})'
                        text_content += f'<div>→ {_esc(link["type"])} {item_ref}</div>'
                    for link in incoming:
                        if link['project_id'] == project.pk:
                            item_ref = f'<a href="#item-{link["item_pk"]}" style="color:#284E8A"><strong>{_esc(link["item_id"])}</strong></a>'
                        else:
                            item_ref = f'<strong>{_esc(link["item_id"])}</strong> ({_esc(link["project"])})'
                        text_content += f'<div>← {_esc(link["type"])} {item_ref}</div>'
                    text_content += '</div>'

            html += f'<tr id="{item_anchor}"><td>{_esc(item.item_id)}</td>'
            html += f'<td>{text_content}</td>'
            html += f'<td><span class="badge {status_class}">{_esc(item.get_status_display())}</span></td>'
            html += f'<td class="{priority_class}">{_esc(item.get_priority_display())}</td>'
            if include_assignee:
                html += f'<td>{_esc(str(item.assignee) if item.assignee else "")}</td>'
            html += '</tr>'

            if options.get('include_comments'):
                comments = list(item.comments.all())
                if comments:
                    colspan = 3 + (1 if include_assignee else 0)
                    html += f'<tr><td></td><td colspan="{colspan}">'
                    for comment in comments:
                        html += f'<div class="comment"><strong>{_esc(str(comment.author))}</strong>: {_esc(comment.text)}</div>'
                    html += '</td></tr>'

        html += '</tbody></table>'

    html += '</body></html>'
    return html


def _export_docx(project, sections, items_by_section, options):
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Inches, Mm, Pt, Emu
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        return _export_docx_fallback(project, sections, items_by_section, options)

    doc = Document()

    settings_el = doc.settings.element
    update_fields = OxmlElement('w:updateFields')
    update_fields.set(qn('w:val'), 'true')
    settings_el.append(update_fields)

    paper, orientation = page_setup(options)
    page_width, page_height = (Mm(side) for side in PAPER_SIZES_MM[paper])
    if orientation == 'landscape':
        page_width, page_height = page_height, page_width
    for section_doc in doc.sections:
        section_doc.orientation = WD_ORIENT.LANDSCAPE if orientation == 'landscape' else WD_ORIENT.PORTRAIT
        section_doc.page_width = page_width
        section_doc.page_height = page_height
        section_doc.top_margin = Emu(914400)
        section_doc.bottom_margin = Emu(914400)
        section_doc.left_margin = Emu(914400)
        section_doc.right_margin = Emu(914400)

    for _blank in range(6):
        doc.add_paragraph()

    title_para = doc.add_heading(project.name, level=0)
    title_para.alignment = 1

    if project.description:
        desc_para = doc.add_paragraph(project.description)
        desc_para.alignment = 1

    version_number = _get_latest_version_number(project)
    if version_number:
        ver_para = doc.add_paragraph()
        ver_para.alignment = 1
        ver_run = ver_para.add_run(_('Version %(number)s') % {'number': version_number})
        ver_run.font.size = Pt(12)
        ver_run.font.color.rgb = _docx_rgb(0x66, 0x66, 0x66)

    if options.get('include_authors'):
        authors = _get_authors(project)
        if authors:
            authors_para = doc.add_paragraph()
            authors_para.alignment = 1
            run = authors_para.add_run(_authors_line(authors))
            run.font.size = Pt(10)
            run.font.color.rgb = _docx_rgb(0x66, 0x66, 0x66)

    _add_section_break(doc, start_type='nextPage', show_page_number=False)

    if options.get('include_toc'):
        toc_heading = doc.add_heading(_('Table of Contents'), level=1)
        toc_heading.alignment = 1
        _add_toc_field(doc, sections)
        _add_section_break(doc, start_type='nextPage', show_page_number=False)

    current_section = doc.sections[-1]
    _set_page_number_footer(
        current_section, start_at=1,
        org_name=options.get('_org_name', ''),
        year=options.get('_year'),
    )

    user_names = options.get('_user_names', {})
    parent_ids = {s.parent_id for s in sections if s.parent_id}
    include_assignee = bool(options.get('include_assignee'))
    page_content_emu = int(current_section.page_width) - int(current_section.left_margin) - int(current_section.right_margin)
    col_id_twips = 1150
    col_status_twips = 1150
    col_priority_twips = 1010
    col_assignee_twips = 1300
    export_attrs = _get_export_attributes(project, options)
    total_twips = page_content_emu * 1440 // 914400
    fixed_twips = col_id_twips + col_status_twips + col_priority_twips
    if include_assignee:
        fixed_twips += col_assignee_twips
    col_text_twips = total_twips - fixed_twips
    img_max_width_emu = (col_text_twips * 914400 // 1440) - int(Inches(0.3))

    for section in sections:
        level = min(section.number.count('.') + 1, 9)
        heading = doc.add_heading('', level=level)
        bookmark_name = f'section_{section.pk}'
        _add_bookmark(heading, bookmark_name, f'{section.number} {section.title}')

        section_items = items_by_section.get(section.pk, [])
        if not section_items:
            if _shows_empty_note(section, items_by_section, parent_ids):
                doc.add_paragraph(_('No items in this section.'), style='Body Text')
            continue

        cols = [_('ID'), _('Text'), _('Status'), _('Priority')]
        col_widths_twips = [col_id_twips, col_text_twips, col_status_twips, col_priority_twips]
        if include_assignee:
            cols.append(_('Assignee'))
            col_widths_twips.append(col_assignee_twips)
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = 'Table Grid'
        table.autofit = False
        _set_table_borders(table, color='D1D5DB', size='4')
        _set_table_grid(table, total_twips, col_widths_twips)

        header_row = table.rows[0]
        _mark_header_row(header_row)
        for i, col_name in enumerate(cols):
            cell = header_row.cells[i]
            _set_cell_width(cell, col_widths_twips[i])
            cell.text = col_name
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)

        for item in section_items:
            row = table.add_row()
            _keep_row_together(row)
            for ci in range(len(cols)):
                _set_cell_width(row.cells[ci], col_widths_twips[ci])

            id_cell = row.cells[0]
            _reset_cell(id_cell)
            bookmark_name = f'item_{item.pk}'
            id_run = _add_bookmark(id_cell.paragraphs[0], bookmark_name, item.item_id)
            id_run.font.size = Pt(9)

            text_cell = row.cells[1]
            _html_to_docx_cell(text_cell, item.text or '', col_text_twips - 120)

            prepared = _prepare_image(item.image)
            if prepared:
                data, _mime, width_px, height_px = prepared
                pic_width = _fit_image_width_emu(width_px, height_px, img_max_width_emu, IMAGE_MAX_HEIGHT_EMU)
                p = text_cell.add_paragraph()
                run = p.add_run()
                run.add_picture(io.BytesIO(data), width=Emu(pic_width))
                if item.image_caption:
                    cap_p = text_cell.add_paragraph()
                    cap_run = cap_p.add_run(item.image_caption)
                    cap_run.italic = True
                    cap_run.font.size = Pt(8)

            attr_lines = _attribute_lines(item, export_attrs, user_names)
            if attr_lines:
                p = text_cell.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                for n, (name, value) in enumerate(attr_lines):
                    if n:
                        sep_run = p.add_run(' · ')
                        sep_run.font.size = Pt(8)
                        sep_run.font.color.rgb = _docx_rgb(0x99, 0x99, 0x99)
                    name_run = p.add_run(f'{name}: ')
                    name_run.bold = True
                    name_run.font.size = Pt(8)
                    name_run.font.color.rgb = _docx_rgb(0x33, 0x33, 0x33)
                    value_run = p.add_run(value)
                    value_run.font.size = Pt(8)
                    value_run.font.color.rgb = _docx_rgb(0x55, 0x55, 0x55)

            if options.get('include_links'):
                outgoing, incoming = _get_item_links(item)
                for direction, links in (('→', outgoing), ('←', incoming)):
                    for link in links:
                        p = text_cell.add_paragraph()
                        p.paragraph_format.space_after = Pt(0)
                        arrow_run = p.add_run(f'{direction} {link["type"]} ')
                        arrow_run.font.size = Pt(8)
                        arrow_run.font.color.rgb = _docx_rgb(0x55, 0x55, 0x55)
                        if link['project_id'] == project.pk:
                            _add_hyperlink_to_bookmark(p, f'item_{link["item_pk"]}', link['item_id'], Pt(8))
                        else:
                            ref_run = p.add_run(f'{link["item_id"]} ({link["project"]})')
                            ref_run.font.size = Pt(8)
                            ref_run.italic = True

            status_cell = row.cells[2]
            _reset_cell(status_cell)
            status_run = status_cell.paragraphs[0].add_run(item.get_status_display())
            status_run.font.size = Pt(8)
            status_run.font.color.rgb = _docx_rgb(*STATUS_COLORS.get(item.status, (0x22, 0x22, 0x22)))

            priority_cell = row.cells[3]
            _reset_cell(priority_cell)
            priority_run = priority_cell.paragraphs[0].add_run(item.get_priority_display())
            priority_run.font.size = Pt(8)
            priority_run.font.color.rgb = _docx_rgb(*PRIORITY_COLORS.get(item.priority, (0x22, 0x22, 0x22)))
            priority_run.bold = item.priority == 'high'

            if include_assignee:
                assignee_cell = row.cells[4]
                _reset_cell(assignee_cell)
                assignee_run = assignee_cell.paragraphs[0].add_run(
                    str(item.assignee) if item.assignee else ''
                )
                assignee_run.font.size = Pt(8)

            if options.get('include_comments'):
                comments = list(item.comments.all())
                if comments:
                    comment_row = table.add_row()
                    _keep_row_together(comment_row)
                    for ci in range(len(cols)):
                        _set_cell_width(comment_row.cells[ci], col_widths_twips[ci])
                    merged = comment_row.cells[1].merge(comment_row.cells[len(cols) - 1])
                    _reset_cell(merged)
                    _shade_cell(merged, 'F8F9FA')
                    for n, comment in enumerate(comments):
                        p = merged.paragraphs[0] if n == 0 else merged.add_paragraph()
                        p.paragraph_format.space_after = Pt(2)
                        run_author = p.add_run(f'{comment.author}: ')
                        run_author.bold = True
                        run_author.font.size = Pt(8)
                        run_text = p.add_run(comment.text)
                        run_text.font.size = Pt(8)

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    docx_buffer.seek(0)

    response = HttpResponse(
        docx_buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="{project.slug}.docx"'
    return response


def _html_to_docx_cell(cell, html, width_twips=None):
    from html.parser import HTMLParser
    from docx.shared import Pt

    if not html or not html.strip():
        cell.text = ''
        return

    class DocxHTMLParser(HTMLParser):
        def __init__(self, cell):
            super().__init__()
            self.cell = cell
            self.current_para = None
            self.current_para_empty = False
            self.current_run_text = ''
            self.bold = False
            self.italic = False
            self.underline = False
            self.strikethrough = False
            self.in_list = None
            self.list_counter = 0
            self.in_table = False
            self.table_rows = []
            self.current_table_row = []
            self.current_cell_text = ''
            self.current_cell_spans = (1, 1)
            self.in_th = False
            self.heading_level = 0
            if cell.paragraphs and cell.paragraphs[0].text == '':
                cell.paragraphs[0].clear()
                self._use_para(cell.paragraphs[0])

        def _use_para(self, paragraph):
            self.current_para = paragraph
            self.current_para_empty = True
            paragraph.paragraph_format.space_after = Pt(2)

        def _flush_run(self):
            if self.current_run_text and self.current_para is not None:
                run = self.current_para.add_run(self.current_run_text)
                run.font.size = Pt(9)
                if self.bold or self.heading_level:
                    run.bold = True
                if self.italic:
                    run.italic = True
                if self.underline:
                    run.underline = True
                if self.strikethrough:
                    run.font.strike = True
                if self.heading_level == 1:
                    run.font.size = Pt(14)
                elif self.heading_level == 2:
                    run.font.size = Pt(12)
                elif self.heading_level == 3:
                    run.font.size = Pt(10)
                self.current_para_empty = False
            self.current_run_text = ''

        def _new_para(self):
            self._flush_run()
            if self.current_para is not None and self.current_para_empty:
                return self.current_para
            self._use_para(self.cell.add_paragraph())
            return self.current_para

        def handle_starttag(self, tag, attrs):
            if tag == 'p':
                self._new_para()
            elif tag == 'br':
                self._flush_run()
                if self.current_para:
                    self.current_para.add_run('\n')
            elif tag == 'strong' or tag == 'b':
                self._flush_run()
                self.bold = True
            elif tag == 'em' or tag == 'i':
                self._flush_run()
                self.italic = True
            elif tag == 'u':
                self._flush_run()
                self.underline = True
            elif tag == 's':
                self._flush_run()
                self.strikethrough = True
            elif tag in ('h1', 'h2', 'h3'):
                self._new_para()
                self.heading_level = int(tag[1])
            elif tag in ('ul', 'ol'):
                self.in_list = tag
                self.list_counter = 0
            elif tag == 'li':
                self._new_para()
                self.list_counter += 1
                if self.in_list == 'ol':
                    self.current_run_text = f'{self.list_counter}. '
                else:
                    self.current_run_text = '• '
            elif tag == 'table':
                self._flush_run()
                self.in_table = True
                self.table_rows = []
            elif tag == 'tr':
                self.current_table_row = []
            elif tag in ('td', 'th'):
                attributes = dict(attrs)
                self.current_cell_text = ''
                self.current_cell_spans = (_span(attributes.get('rowspan')), _span(attributes.get('colspan')))
                self.in_th = (tag == 'th')

        def handle_endtag(self, tag):
            if tag == 'p':
                self._flush_run()
            elif tag == 'strong' or tag == 'b':
                self._flush_run()
                self.bold = False
            elif tag == 'em' or tag == 'i':
                self._flush_run()
                self.italic = False
            elif tag == 'u':
                self._flush_run()
                self.underline = False
            elif tag == 's':
                self._flush_run()
                self.strikethrough = False
            elif tag in ('h1', 'h2', 'h3'):
                self._flush_run()
                self.heading_level = 0
            elif tag in ('ul', 'ol'):
                self.in_list = None
                self.list_counter = 0
            elif tag == 'li':
                self._flush_run()
            elif tag in ('td', 'th'):
                rowspan, colspan = self.current_cell_spans
                self.current_table_row.append({
                    'text': self.current_cell_text,
                    'bold': self.in_th,
                    'rowspan': rowspan,
                    'colspan': colspan,
                })
            elif tag == 'tr':
                if self.current_table_row:
                    self.table_rows.append(self.current_table_row)
            elif tag == 'table':
                self.in_table = False
                self._render_table()

        def handle_data(self, data):
            if self.in_table and not self.in_list:
                self.current_cell_text += data
            else:
                self.current_run_text += data

        def _render_table(self):
            placed, n_rows, n_cols = _table_grid(self.table_rows)
            self.table_rows = []
            if not placed:
                return
            tbl = self.cell.add_table(rows=n_rows, cols=n_cols)
            tbl.autofit = False
            inner_width = width_twips or 4000
            col_width = inner_width // n_cols
            _set_table_grid(tbl, col_width * n_cols, [col_width] * n_cols)
            _set_table_borders(tbl, color='D1D5DB', size='4')
            for r, c, rowspan, colspan, cell_data in placed:
                target = tbl.cell(r, c)
                if rowspan > 1 or colspan > 1:
                    target = target.merge(tbl.cell(r + rowspan - 1, c + colspan - 1))
                _reset_cell(target)
                _set_cell_width(target, col_width * colspan)
                paragraph = target.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(cell_data['text'].strip())
                run.font.size = Pt(8)
                if cell_data.get('bold'):
                    run.bold = True
            self._use_para(self.cell.paragraphs[-1])

    cell.text = ''
    parser = DocxHTMLParser(cell)
    parser.feed(html)
    parser._flush_run()


def _span(value):
    try:
        return max(1, int(value))
    except (TypeError, ValueError):
        return 1


def _table_grid(rows):
    occupied = set()
    placed = []
    n_rows = 0
    n_cols = 0
    for r, row in enumerate(rows):
        c = 0
        for cell_data in row:
            while (r, c) in occupied:
                c += 1
            rowspan = cell_data['rowspan']
            colspan = cell_data['colspan']
            for dr in range(rowspan):
                for dc in range(colspan):
                    occupied.add((r + dr, c + dc))
            placed.append((r, c, rowspan, colspan, cell_data))
            n_rows = max(n_rows, r + rowspan)
            n_cols = max(n_cols, c + colspan)
            c += colspan
    return placed, n_rows, n_cols


def _reset_cell(cell):
    for paragraph in cell.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    cell.paragraphs[0].clear()


def _shade_cell(cell, fill):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _mark_header_row(row):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    trPr.append(header)


def _keep_row_together(row):
    from docx.oxml import OxmlElement
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))


def _set_table_grid(table, total_twips, col_widths_twips):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblLayout = tblPr.find(qn('w:tblLayout'))
    if tblLayout is None:
        tblLayout = OxmlElement('w:tblLayout')
        tblPr.append(tblLayout)
    tblLayout.set(qn('w:type'), 'fixed')
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(total_twips))
    tblW.set(qn('w:type'), 'dxa')
    old_grid = tbl.find(qn('w:tblGrid'))
    if old_grid is not None:
        tbl.remove(old_grid)
    tblGrid = OxmlElement('w:tblGrid')
    for w in col_widths_twips:
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), str(w))
        tblGrid.append(gridCol)
    tbl.insert(tbl.index(tblPr) + 1, tblGrid)


def _docx_rgb(r, g, b):
    from docx.shared import RGBColor
    return RGBColor(r, g, b)


def _add_bookmark(paragraph, bookmark_name, text):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tag = OxmlElement('w:bookmarkStart')
    tag.set(qn('w:id'), str(id(bookmark_name) % 2147483647))
    tag.set(qn('w:name'), bookmark_name)
    paragraph._p.append(tag)

    run = paragraph.add_run(text)

    tag_end = OxmlElement('w:bookmarkEnd')
    tag_end.set(qn('w:id'), str(id(bookmark_name) % 2147483647))
    paragraph._p.append(tag_end)

    return run


def _add_hyperlink_to_bookmark(paragraph, bookmark_name, text, font_size=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), bookmark_name)

    run_el = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '284E8A')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    if font_size:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(font_size.pt * 2)))
        rPr.append(sz)
    run_el.append(rPr)

    text_el = OxmlElement('w:t')
    text_el.set(qn('xml:space'), 'preserve')
    text_el.text = text
    run_el.append(text_el)

    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def _set_table_borders(table, color='D1D5DB', size='4'):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    existing = tblPr.find(qn('w:tblBorders'))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def _set_cell_width(cell, width_twips):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(width_twips))
    tcW.set(qn('w:type'), 'dxa')


def _set_vertical_alignment(section, align='center'):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    sect_pr = section._sectPr
    vAlign = sect_pr.find(qn('w:vAlign'))
    if vAlign is None:
        vAlign = OxmlElement('w:vAlign')
        sect_pr.append(vAlign)
    vAlign.set(qn('w:val'), align)


def _add_section_break(doc, start_type='nextPage', show_page_number=True):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    new_section = doc.add_section()
    sect_pr = new_section._sectPr

    type_el = sect_pr.find(qn('w:type'))
    if type_el is None:
        type_el = OxmlElement('w:type')
        sect_pr.insert(0, type_el)
    type_el.set(qn('w:val'), start_type)

    if not show_page_number:
        for footer_ref in sect_pr.findall(qn('w:footerReference')):
            sect_pr.remove(footer_ref)


def _add_field(paragraph, instruction, font_size, color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    for part in ('begin', 'instr', 'end'):
        run = paragraph.add_run()
        run.font.size = font_size
        run.font.color.rgb = color
        if part == 'instr':
            instr = OxmlElement('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = f' {instruction} '
            run._r.append(instr)
        else:
            fld = OxmlElement('w:fldChar')
            fld.set(qn('w:fldCharType'), part)
            run._r.append(fld)


def _set_page_number_footer(section, start_at=1, org_name='', year=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    if year is None:
        year = date.today().year

    footer_left = f'{org_name} © {year}' if org_name else f'© {year}'
    grey = _docx_rgb(0x99, 0x99, 0x99)
    size = Pt(8)

    sect_pr = section._sectPr
    pgNumType = sect_pr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sect_pr.append(pgNumType)
    pgNumType.set(qn('w:start'), str(start_at))

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.clear()

    content_width = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
    twips = content_width // 635
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'right')
    tab.set(qn('w:pos'), str(twips))
    tabs.append(tab)
    pPr.append(tabs)

    for text in (footer_left, '\t'):
        run = p.add_run(text)
        run.font.size = size
        run.font.color.rgb = grey
    for kind, value in _page_label_parts():
        if kind == 'field':
            _add_field(p, PAGE_FIELD_CODES[value], size, grey)
        else:
            run = p.add_run(value)
            run.font.size = size
            run.font.color.rgb = grey


def _add_toc_field(doc, sections=None):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    fldChar.set(qn('w:dirty'), 'true')
    run._r.append(fldChar)

    run2 = paragraph.add_run()
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' TOC \\o "1-3" \\h \\z \\u '
    run2._r.append(instrText)

    run3 = paragraph.add_run()
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    run3._r.append(fldChar2)

    if sections:
        for section in sections:
            toc_para = doc.add_paragraph()
            toc_para.style = doc.styles['Normal']
            depth = section.number.count('.')
            indent_cm = depth * 1.0
            if indent_cm:
                pPr = toc_para._p.get_or_add_pPr()
                ind = OxmlElement('w:ind')
                ind.set(qn('w:left'), str(int(indent_cm * 567)))
                pPr.append(ind)
            font_size = Pt(11) if depth == 0 else Pt(10)
            bookmark_name = f'section_{section.pk}'
            _add_hyperlink_to_bookmark(
                toc_para, bookmark_name,
                f'{section.number}  {section.title}',
                font_size=font_size,
            )
    else:
        fallback_run = paragraph.add_run(_('[Table of contents will appear here]'))
        fallback_run.italic = True

    end_para = doc.add_paragraph()
    end_run = end_para.add_run()
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    end_run._r.append(fldChar3)


def _esc(text):
    if not text:
        return ''
    return (
        str(text)
        .replace('&', '&amp;')
        .replace('<', '&lt;')
        .replace('>', '&gt;')
        .replace('"', '&quot;')
    )


def _export_pdf_fallback(project, sections, items_by_section, options):
    response = HttpResponse(content_type='text/plain')
    response['Content-Disposition'] = f'attachment; filename="{project.slug}-export.txt"'
    lines = [f'{project.name}\n{"=" * len(project.name)}\n']
    if project.description:
        lines.append(f'{project.description}\n')
    for section in sections:
        lines.append(f'\n{section.number} {section.title}')
        for item in items_by_section.get(section.pk, []):
            lines.append(f'  {item.item_id}: {item.text} [{item.status}] [{item.priority}]')
    response.write('\n'.join(lines).encode('utf-8'))
    return response


def _export_docx_fallback(project, sections, items_by_section, options):
    response = HttpResponse(content_type='text/plain')
    response['Content-Disposition'] = f'attachment; filename="{project.slug}-export.txt"'
    lines = [f'{project.name}\n']
    if project.description:
        lines.append(f'{project.description}\n')
    for section in sections:
        lines.append(f'\n{section.number} {section.title}')
        for item in items_by_section.get(section.pk, []):
            lines.append(f'  {item.item_id}: {item.text} [{item.status}]')
    response.write('\n'.join(lines).encode('utf-8'))
    return response
