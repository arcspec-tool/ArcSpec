import base64
import io
import re
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.test import TestCase
from django.urls import reverse
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Inches, Mm
from PIL import Image, ImageDraw

from apps.accounts.factories import UserFactory
from apps.projects.export import IMAGE_MAX_HEIGHT_CM
from apps.projects.factories import (
    CommentFactory,
    ItemAttributeValueFactory,
    ItemFactory,
    ProjectAttributeFactory,
    ProjectFactory,
    ProjectMembershipFactory,
    SectionFactory,
    TraceabilityLinkFactory,
)


def docx_text(document):
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
                for inner in cell.tables:
                    for inner_row in inner.rows:
                        for inner_cell in inner_row.cells:
                            parts.extend(p.text for p in inner_cell.paragraphs)
    return '\n'.join(parts)


class ExportOptionsTest(TestCase):

    def setUp(self):
        self.user = UserFactory()
        self.project = ProjectFactory()
        ProjectMembershipFactory(project=self.project, user=self.user, role='editor')
        section = SectionFactory(project=self.project, number='1', title='Intro')
        self.item = ItemFactory(project=self.project, section=section, item_id='REQ-001', text='<p>First</p>', created_by=self.user, order=1, priority='high')
        other = ItemFactory(project=self.project, section=section, item_id='REQ-002', text='<p>Second</p>', created_by=self.user, order=2)
        self.attribute = ProjectAttributeFactory(project=self.project, name='Risk level')
        ItemAttributeValueFactory(item=self.item, attribute=self.attribute, value='High risk')
        CommentFactory(item=self.item, author=self.user, text='A remark in the discussion')
        TraceabilityLinkFactory(source_item=self.item, target_item=other, link_type='verified-by', created_by=self.user)
        self.client.force_login(self.user)
        self.url = reverse('project_export', kwargs={'slug': self.project.slug})
        self.everything = {
            'include_toc': 'on', 'include_attributes': 'on', 'export_attr_ids': [self.attribute.pk],
            'include_comments': 'on', 'include_links': 'on',
        }

    def pdf_html(self, **options):
        with patch('weasyprint.HTML') as html_class:
            response = self.client.post(self.url, {'format': 'pdf', **options})
        self.assertEqual(response.status_code, 200)
        return html_class.call_args.kwargs['string']

    def docx(self, **options):
        response = self.client.post(self.url, {'format': 'docx', **options})
        self.assertEqual(response.status_code, 200)
        return Document(io.BytesIO(response.content))

    def test_the_pdf_page_follows_paper_size_and_orientation(self):
        self.assertIn('size: LETTER landscape;', self.pdf_html(paper_size='letter', orientation='landscape'))
        self.assertIn('size: LEGAL portrait;', self.pdf_html(paper_size='legal'))
        self.assertIn('size: A4 portrait;', self.pdf_html())
        self.assertIn('size: A4 portrait;', self.pdf_html(paper_size='a0', orientation='sideways'))

    def test_each_pdf_include_option_adds_its_content(self):
        html = self.pdf_html(**self.everything)
        for text in ('Table of Contents', 'Risk level', 'High risk', 'A remark in the discussion', 'Verified by'):
            self.assertIn(text, html)
        html = self.pdf_html()
        for text in ('Table of Contents', 'Risk level', 'High risk', 'A remark in the discussion', 'Verified by'):
            self.assertNotIn(text, html)
        self.assertIn('REQ-001', html)

    @staticmethod
    def page_of(section):
        return round(section.page_width / 635), round(section.page_height / 635), section.orientation

    def test_the_docx_page_follows_paper_size_and_orientation(self):
        a4_portrait = (round(Mm(210) / 635), round(Mm(297) / 635), WD_ORIENT.PORTRAIT)
        page = self.docx(paper_size='letter', orientation='landscape').sections[0]
        self.assertEqual(self.page_of(page), (round(Mm(279.4) / 635), round(Mm(215.9) / 635), WD_ORIENT.LANDSCAPE))
        page = self.docx(paper_size='legal').sections[0]
        self.assertEqual(self.page_of(page), (round(Mm(215.9) / 635), round(Mm(355.6) / 635), WD_ORIENT.PORTRAIT))
        for page in self.docx().sections:
            self.assertEqual(self.page_of(page), a4_portrait)
        self.assertEqual(self.page_of(self.docx(paper_size='a0', orientation='sideways').sections[0]), a4_portrait)

    def test_each_docx_include_option_adds_its_content(self):
        text = docx_text(self.docx(**self.everything))
        for part in ('Table of Contents', 'Risk level', 'High risk', 'A remark in the discussion', 'Verified by'):
            self.assertIn(part, text)
        text = docx_text(self.docx())
        for part in ('Table of Contents', 'Risk level', 'High risk', 'A remark in the discussion', 'Verified by'):
            self.assertNotIn(part, text)
        self.assertIn('REQ-001', text)

    def test_attributes_are_listed_under_the_text_and_never_as_columns(self):
        html = self.pdf_html(**self.everything, include_assignee='on')
        self.assertIn('<th style="width:70px">ID</th><th>Text</th><th style="width:70px">Status</th><th style="width:60px">Priority</th><th style="width:90px">Assignee</th></tr>', html)
        self.assertNotIn('<th style="width:80px">Risk level</th>', html)
        self.assertIn('<div class="attrs"><strong>Risk level:</strong> High risk</div>', html)
        table = self.docx(**self.everything, include_assignee='on').tables[0]
        self.assertEqual([cell.text for cell in table.rows[0].cells], ['ID', 'Text', 'Status', 'Priority', 'Assignee'])
        self.assertEqual(len(self.docx(**self.everything).tables[0].rows[0].cells), 4)
        self.assertIn('Risk level: High risk', docx_text(self.docx(**self.everything)))

    def test_docx_comments_sit_in_the_item_row(self):
        document = self.docx(include_comments='on')
        table = document.tables[0]
        self.assertEqual(len(table.rows), 4)
        comment_row = table.rows[2]
        self.assertEqual(comment_row.cells[0].text, '')
        self.assertIn('A remark in the discussion', comment_row.cells[1].text)
        self.assertNotIn('A remark in the discussion', '\n'.join(p.text for p in document.paragraphs))

    def test_docx_text_cell_starts_with_the_text(self):
        table = self.docx().tables[0]
        self.assertEqual(table.rows[1].cells[1].paragraphs[0].text, 'First')
        self.assertEqual(table.rows[1].cells[0].text, 'REQ-001')
        status_run = table.rows[1].cells[2].paragraphs[0].runs[0]
        priority_run = table.rows[1].cells[3].paragraphs[0].runs[0]
        self.assertEqual((status_run.text, str(status_run.font.color.rgb)), ('Draft', '475569'))
        self.assertEqual((priority_run.text, str(priority_run.font.color.rgb), priority_run.bold), ('High', 'DC2626', True))

    def test_docx_footer_counts_pages_and_rows_stay_together(self):
        document = self.docx(include_toc='on')
        footer_xml = document.sections[-1].footer._element.xml
        self.assertIn('PAGE', footer_xml)
        self.assertIn('NUMPAGES', footer_xml)
        self.assertIn(' of ', footer_xml)
        table_xml = document.tables[0]._tbl.xml
        self.assertIn('w:tblHeader', table_xml)
        self.assertIn('w:cantSplit', table_xml)


class ExportLayoutTests(TestCase):

    def setUp(self):
        self.user = UserFactory(display_name='Reviewer One', first_name='', last_name='')
        self.project = ProjectFactory()
        ProjectMembershipFactory(project=self.project, user=self.user, role='editor')
        self.parent = SectionFactory(project=self.project, number='1', title='Parent', order=1)
        self.child = SectionFactory(project=self.project, number='1.1', title='Child', parent=self.parent, order=2)
        self.grandchild = SectionFactory(project=self.project, number='1.1.1', title='Grandchild', parent=self.child, order=3)
        SectionFactory(project=self.project, number='2', title='Empty', order=4)
        self.item = ItemFactory(
            project=self.project, section=self.grandchild, item_id='REQ-001', created_by=self.user,
            text='<p>Lead</p><table><tr><td rowspan="2">A</td><td>B</td></tr><tr><td>C</td></tr>'
                 '<tr><td colspan="2">D</td></tr></table><p>Tail</p>',
        )
        self.reviewer = ProjectAttributeFactory(project=self.project, name='Reviewer', field_type='user')
        self.safety = ProjectAttributeFactory(project=self.project, name='Safety', field_type='boolean')
        self.blank = ProjectAttributeFactory(project=self.project, name='Blank', field_type='text')
        ItemAttributeValueFactory(item=self.item, attribute=self.reviewer, value=str(self.user.pk))
        ItemAttributeValueFactory(item=self.item, attribute=self.safety, value='true')
        ItemAttributeValueFactory(item=self.item, attribute=self.blank, value='')
        self.client.force_login(self.user)
        self.url = reverse('project_export', kwargs={'slug': self.project.slug})
        self.options = {
            'include_toc': 'on', 'include_authors': 'on', 'include_attributes': 'on',
            'export_attr_ids': [self.reviewer.pk, self.safety.pk, self.blank.pk],
        }

    def pdf_html(self):
        with patch('weasyprint.HTML') as html_class:
            response = self.client.post(self.url, {'format': 'pdf', **self.options})
        self.assertEqual(response.status_code, 200)
        return html_class.call_args.kwargs['string']

    def docx(self):
        response = self.client.post(self.url, {'format': 'docx', **self.options})
        self.assertEqual(response.status_code, 200)
        return Document(io.BytesIO(response.content))

    def test_only_a_leaf_section_without_items_gets_the_empty_note(self):
        html = self.pdf_html()
        self.assertEqual(html.count('No items in this section.'), 1)
        self.assertIn('Empty</h2><p><em>No items in this section.</em></p>', html)
        text = docx_text(self.docx())
        self.assertEqual(text.count('No items in this section.'), 1)

    def test_pdf_headings_and_toc_follow_the_section_depth(self):
        html = self.pdf_html()
        self.assertIn(f'<h2 id="section-{self.parent.pk}">1 Parent</h2>', html)
        self.assertIn(f'<h3 id="section-{self.child.pk}">1.1 Child</h3>', html)
        self.assertIn(f'<h4 id="section-{self.grandchild.pk}">1.1.1 Grandchild</h4>', html)
        self.assertIn(f'<li style="padding-left:40px"><a href="#section-{self.grandchild.pk}">1.1.1 Grandchild</a></li>', html)
        self.assertIn(f'<li style="padding-left:0px"><a href="#section-{self.parent.pk}">1 Parent</a></li>', html)

    def test_user_and_boolean_attributes_export_their_display_values(self):
        html = self.pdf_html()
        self.assertIn('<strong>Reviewer:</strong> Reviewer One', html)
        self.assertIn('<strong>Safety:</strong> Yes', html)
        self.assertNotIn(f'<strong>Reviewer:</strong> {self.user.pk}', html)
        self.assertNotIn('Blank:', html)
        text = docx_text(self.docx())
        self.assertIn('Reviewer: Reviewer One', text)
        self.assertIn('Safety: Yes', text)
        self.assertNotIn('Blank:', text)

    def test_authors_use_the_display_name_rather_than_the_email(self):
        html = self.pdf_html()
        self.assertIn(f'Authors: {self.project.created_by}, Reviewer One', html)
        self.assertNotIn(self.user.email, html)
        self.assertIn(f'Authors: {self.project.created_by}, Reviewer One', docx_text(self.docx()))

    def test_docx_inner_table_keeps_the_grid_of_spanned_cells(self):
        cell = self.docx().tables[0].rows[1].cells[1]
        self.assertEqual(cell.paragraphs[0].text, 'Lead')
        inner = cell.tables[0]
        self.assertEqual((len(inner.rows), len(inner.columns)), (3, 2))
        self.assertEqual(inner.cell(0, 0).text, 'A')
        self.assertEqual(inner.cell(0, 1).text, 'B')
        self.assertEqual(inner.cell(1, 0).text, 'A')
        self.assertEqual(inner.cell(1, 1).text, 'C')
        self.assertEqual(inner.cell(2, 0).text, 'D')
        self.assertEqual(inner.cell(2, 1).text, 'D')
        self.assertEqual([p.text for p in cell.paragraphs], ['Lead', 'Tail', 'Reviewer: Reviewer One · Safety: Yes'])


class ImageExportTests(TestCase):

    def setUp(self):
        self.user = UserFactory()
        self.project = ProjectFactory()
        ProjectMembershipFactory(project=self.project, user=self.user, role='editor')
        self.section = SectionFactory(project=self.project, number='1', title='Pictures')
        self.items = {}
        for n, (name, data) in enumerate(self.pictures().items(), start=1):
            item = ItemFactory(project=self.project, section=self.section, item_id=f'REQ-{n:03d}', text=f'<p>{name}</p>', created_by=self.user, order=n)
            item.image.save(name, ContentFile(data), save=True)
            self.items[name] = item
        self.client.force_login(self.user)
        self.url = reverse('project_export', kwargs={'slug': self.project.slug})

    @staticmethod
    def encode(image, fmt, **options):
        buffer = io.BytesIO()
        image.save(buffer, format=fmt, **options)
        return buffer.getvalue()

    def pictures(self):
        upright = Image.new('RGB', (120, 160), 'white')
        ImageDraw.Draw(upright).rectangle([0, 0, 119, 20], fill='red')
        exif = Image.Exif()
        exif[0x0112] = 6
        return {
            'rotated.jpg': self.encode(upright.rotate(90, expand=True), 'JPEG', exif=exif.tobytes()),
            'photo.webp': self.encode(Image.new('RGB', (100, 50), 'blue'), 'WEBP'),
            'deep.png': self.encode(Image.new('I;16', (40, 30), 40000), 'PNG'),
            'print.jpg': self.encode(Image.new('CMYK', (50, 40), (0, 0, 0, 0)), 'JPEG'),
            'tall.png': self.encode(Image.new('RGB', (100, 3000), 'green'), 'PNG'),
            'plain.png': self.encode(Image.new('RGBA', (30, 20), (0, 0, 255, 128)), 'PNG'),
        }

    def pdf_images(self):
        with patch('weasyprint.HTML') as html_class:
            response = self.client.post(self.url, {'format': 'pdf'})
        self.assertEqual(response.status_code, 200)
        html = html_class.call_args.kwargs['string']
        self.assertIn(f'max-height: {IMAGE_MAX_HEIGHT_CM}cm', html)
        images = {}
        for item_id, uri in re.findall(r'<td>(REQ-\d{3})</td><td>.*?<img class="item-image" src="(data:[^"]+)">', html, re.S):
            mime, payload = uri[5:].split(';base64,')
            images[item_id] = (mime, Image.open(io.BytesIO(base64.b64decode(payload))))
        return images

    def test_pdf_pictures_are_upright_eight_bit_and_rgb(self):
        images = self.pdf_images()
        self.assertEqual(len(images), 6)
        mime, rotated = images[self.items['rotated.jpg'].item_id]
        self.assertEqual((mime, rotated.size), ('image/jpeg', (120, 160)))
        self.assertEqual(rotated.getpixel((60, 5))[0] > 200 and rotated.getpixel((60, 5))[2] < 80, True)
        mime, webp = images[self.items['photo.webp'].item_id]
        self.assertEqual((mime, webp.format, webp.size), ('image/jpeg', 'JPEG', (100, 50)))
        mime, deep = images[self.items['deep.png'].item_id]
        self.assertEqual((deep.mode, deep.size), ('L', (40, 30)))
        self.assertGreater(deep.getpixel((1, 1)), 120)
        mime, cmyk = images[self.items['print.jpg'].item_id]
        self.assertEqual(cmyk.mode, 'RGB')
        mime, plain = images[self.items['plain.png'].item_id]
        self.assertEqual((mime, plain.format, plain.mode), ('image/png', 'PNG', 'RGBA'))

    def test_docx_embeds_every_picture_within_the_page(self):
        response = self.client.post(self.url, {'format': 'docx'})
        self.assertEqual(response.status_code, 200)
        document = Document(io.BytesIO(response.content))
        shapes = list(document.inline_shapes)
        self.assertEqual(len(shapes), 6)
        by_item = {}
        for item_id, shape in zip(sorted(self.items[name].item_id for name in self.items), shapes):
            by_item[item_id] = shape
        rotated = by_item[self.items['rotated.jpg'].item_id]
        self.assertLess(rotated.width, rotated.height)
        tall = by_item[self.items['tall.png'].item_id]
        self.assertLessEqual(tall.height, Cm(IMAGE_MAX_HEIGHT_CM))
        self.assertGreater(tall.height, Cm(IMAGE_MAX_HEIGHT_CM - 0.1))
        self.assertEqual(round(tall.width / tall.height, 2), round(100 / 3000, 2))
        for shape in shapes:
            self.assertLessEqual(shape.width, Mm(210) - Inches(2))
